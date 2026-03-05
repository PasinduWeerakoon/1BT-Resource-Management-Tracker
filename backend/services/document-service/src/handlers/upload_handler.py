"""
Excel Upload Handler
Handles bulk data import from Excel files for projects, employees, and allocations.

Upload format: Base64-encoded Excel file in request body
Each handler validates headers, parses rows, and performs upserts to the database.
"""

import json
import base64
import io
import logging
from datetime import datetime

import openpyxl
from ..utils.database import get_connection
from ..utils.response import success, error

logger = logging.getLogger(__name__)

# ============================================================================
# HELPER FUNCTIONS
# ============================================================================

def _parse_excel_from_event(event):
    """Parse base64-encoded Excel file from Lambda event body"""
    body = event.get('body', '')
    is_base64 = event.get('isBase64Encoded', False)

    if not body:
        raise ValueError("Request body is empty. Send a base64-encoded Excel file.")

    try:
        if isinstance(body, str):
            parsed = json.loads(body)
            file_data = parsed.get('file', parsed.get('data', ''))
        else:
            file_data = body
    except json.JSONDecodeError:
        file_data = body

    if is_base64 or (isinstance(file_data, str) and len(file_data) > 100):
        try:
            # Strip data URI prefix if present (e.g., data:application/...;base64,)
            if isinstance(file_data, str) and ';base64,' in file_data:
                file_data = file_data.split(';base64,')[1]
            decoded = base64.b64decode(file_data)
            return io.BytesIO(decoded)
        except Exception as e:
            raise ValueError(f"Failed to decode base64 file: {str(e)}")

    raise ValueError("Could not parse Excel file from request body. Send base64-encoded data.")


def _parse_date(value):
    """Safely parse a date value from Excel"""
    if value is None or str(value).strip() == '' or str(value).strip().lower() == 'n/a':
        return None
    if isinstance(value, datetime):
        return value.strftime('%Y-%m-%d')
    try:
        return str(value).strip()
    except Exception:
        return None


def _clean_str(value):
    """Clean a string value from Excel"""
    if value is None:
        return None
    return str(value).strip() if str(value).strip() else None


def _clean_int(value):
    """Clean an integer value from Excel"""
    if value is None or str(value).strip() == '' or str(value).strip().lower() == 'n/a':
        return None
    try:
        return int(float(str(value).strip()))
    except (ValueError, TypeError):
        return None


def _clean_decimal(value):
    """Clean a decimal value from Excel"""
    if value is None or str(value).strip() == '' or str(value).strip().lower() == 'n/a':
        return None
    try:
        return float(str(value).strip())
    except (ValueError, TypeError):
        return None


def _normalize_header(header):
    """Normalize Excel header to snake_case for matching"""
    if header is None:
        return ''
    return str(header).strip().lower().replace(' ', '_').replace('-', '_').replace('/', '_')


def _get_headers(ws):
    """Extract and normalize headers from first row"""
    headers = []
    for cell in ws[1]:
        headers.append(_normalize_header(cell.value))
    return headers


# ============================================================================
# PROJECT UPLOAD
# ============================================================================

def upload_projects(event, context):
    """
    Upload projects from Excel file.
    Expected columns: Project Name, Project ID/Code, Project Type, Client Name,
                     Project Start Date, Project End Date, Status, Account Manager,
                     Account Type, Sales Owner, Billing Status, Description
    """
    logger.info("Starting project upload from Excel")

    try:
        file_buffer = _parse_excel_from_event(event)
        wb = openpyxl.load_workbook(file_buffer, read_only=True, data_only=True)
        ws = wb.active

        headers = _get_headers(ws)
        logger.info(f"Found headers: {headers}")

        conn = get_connection()
        conn.autocommit = False
        cursor = conn.cursor()

        results = {'imported': 0, 'updated': 0, 'failed': 0, 'errors': []}

        # Header mapping (flexible matching)
        HEADER_MAP = {
            'project_name': ['project_name', 'name', 'project'],
            'project_code': ['project_code', 'project_id', 'code', 'id'],
            'project_type': ['project_type', 'type'],
            'client_name': ['client_name', 'client', 'account_name'],
            'project_start_date': ['project_start_date', 'start_date', 'start'],
            'project_end_date': ['project_end_date', 'end_date', 'end'],
            'status': ['status', 'project_status'],
            'account_manager': ['account_manager', 'manager', 'am'],
            'account_type': ['account_type', 'account', 'acc_type'],
            'sales_owner': ['sales_owner', 'account_reg_sales_owner', 'sales'],
            'billing_status': ['billing_status', 'billing'],
            'description': ['description', 'desc', 'notes'],
        }

        def find_col(field_name):
            aliases = HEADER_MAP.get(field_name, [field_name])
            for alias in aliases:
                if alias in headers:
                    return headers.index(alias)
            return None

        col_map = {field: find_col(field) for field in HEADER_MAP}

        if col_map['project_name'] is None:
            return error("Required column 'Project Name' not found in Excel headers", 400,
                        {'found_headers': headers})

        try:
            for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
                try:
                    def get_val(field):
                        idx = col_map.get(field)
                        return row[idx] if idx is not None and idx < len(row) else None

                    project_name = _clean_str(get_val('project_name'))
                    if not project_name:
                        continue

                    project_code = _clean_str(get_val('project_code'))
                    project_type = _clean_str(get_val('project_type'))
                    client_name = _clean_str(get_val('client_name'))
                    start_date = _parse_date(get_val('project_start_date'))
                    end_date = _parse_date(get_val('project_end_date'))
                    status = _clean_str(get_val('status')) or 'Active'
                    account_manager_name = _clean_str(get_val('account_manager'))
                    account_type = _clean_str(get_val('account_type')) or 'External'
                    sales_owner = _clean_str(get_val('sales_owner'))
                    billing_status_name = _clean_str(get_val('billing_status'))
                    description = _clean_str(get_val('description'))

                    # Resolve project_type_id
                    project_type_id = None
                    if project_type:
                        cursor.execute("SELECT id FROM project_types WHERE name = %s", (project_type,))
                        pt_row = cursor.fetchone()
                        if pt_row:
                            project_type_id = pt_row['id']

                    # Resolve client_id (create client if needed)
                    client_id = None
                    if client_name:
                        cursor.execute("SELECT id FROM clients WHERE client_name = %s AND deleted_at IS NULL", (client_name,))
                        cl_row = cursor.fetchone()
                        if cl_row:
                            client_id = cl_row['id']
                        else:
                            cursor.execute(
                                "INSERT INTO clients (client_name, is_active) VALUES (%s, TRUE) RETURNING id",
                                (client_name,)
                            )
                            client_id = cursor.fetchone()['id']

                    # Resolve account_manager_id
                    account_manager_id = None
                    if account_manager_name:
                        cursor.execute(
                            "SELECT id FROM employees WHERE name ILIKE %s AND deleted_at IS NULL LIMIT 1",
                            (f"%{account_manager_name}%",)
                        )
                        am_row = cursor.fetchone()
                        if am_row:
                            account_manager_id = am_row['id']

                    # Resolve billing_status_id
                    billing_status_id = None
                    if billing_status_name:
                        cursor.execute("SELECT id FROM billing_statuses WHERE name = %s", (billing_status_name,))
                        bs_row = cursor.fetchone()
                        if bs_row:
                            billing_status_id = bs_row['id']

                    # Get system user for created_by
                    cursor.execute("SELECT id FROM users WHERE username = 'system' LIMIT 1")
                    system_user = cursor.fetchone()
                    created_by = system_user['id'] if system_user else 1

                    # Upsert project
                    cursor.execute("""
                        INSERT INTO projects (
                            project_name, project_code, project_type_id, client_id,
                            project_start_date, project_end_date, status, account_manager_id,
                            account_type, account_reg_sales_owner, billing_status_id,
                            description, created_by
                        ) VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
                        ON CONFLICT (project_name) WHERE deleted_at IS NULL
                        DO UPDATE SET
                            project_code = COALESCE(EXCLUDED.project_code, projects.project_code),
                            project_type_id = COALESCE(EXCLUDED.project_type_id, projects.project_type_id),
                            client_id = COALESCE(EXCLUDED.client_id, projects.client_id),
                            project_start_date = COALESCE(EXCLUDED.project_start_date, projects.project_start_date),
                            project_end_date = COALESCE(EXCLUDED.project_end_date, projects.project_end_date),
                            status = EXCLUDED.status,
                            account_manager_id = COALESCE(EXCLUDED.account_manager_id, projects.account_manager_id),
                            account_type = EXCLUDED.account_type,
                            account_reg_sales_owner = COALESCE(EXCLUDED.account_reg_sales_owner, projects.account_reg_sales_owner),
                            billing_status_id = COALESCE(EXCLUDED.billing_status_id, projects.billing_status_id),
                            description = COALESCE(EXCLUDED.description, projects.description),
                            updated_at = NOW()
                        RETURNING (xmax = 0) AS is_insert
                    """, (
                        project_name, project_code, project_type_id, client_id,
                        start_date, end_date, status, account_manager_id,
                        account_type, sales_owner, billing_status_id,
                        description, created_by
                    ))

                    result_row = cursor.fetchone()
                    if result_row and result_row['is_insert']:
                        results['imported'] += 1
                    else:
                        results['updated'] += 1

                except Exception as row_err:
                    results['failed'] += 1
                    results['errors'].append({'row': row_idx, 'error': str(row_err)})
                    logger.warning(f"Row {row_idx} failed: {str(row_err)}")

            conn.commit()
            logger.info(f"Project upload complete: {results}")
            return success({
                'message': 'Project upload completed',
                'results': results
            })

        except Exception as e:
            conn.rollback()
            raise e

    except ValueError as ve:
        return error(str(ve), 400)
    except Exception as e:
        logger.error(f"Project upload failed: {str(e)}", exc_info=True)
        return error(f"Project upload failed: {str(e)}", 500)


# ============================================================================
# EMPLOYEE UPLOAD
# ============================================================================

def upload_employees(event, context):
    """
    Upload employees from Excel file.
    Expected columns: EPF No, Emp No, Name, Email, Track, Tech Stack, Tier,
                     Designation, Employee Type, University, Joined Date,
                     Date of Birth, Last Increment Date, Last Promotion Date,
                     Status, Is Account Manager, Is External
    """
    logger.info("Starting employee upload from Excel")

    try:
        file_buffer = _parse_excel_from_event(event)
        wb = openpyxl.load_workbook(file_buffer, read_only=True, data_only=True)
        ws = wb.active

        headers = _get_headers(ws)
        logger.info(f"Found headers: {headers}")

        conn = get_connection()
        conn.autocommit = False
        cursor = conn.cursor()

        results = {'imported': 0, 'updated': 0, 'failed': 0, 'errors': []}

        HEADER_MAP = {
            'epf_no': ['epf_no', 'epf', 'epf_number'],
            'emp_no': ['emp_no', 'employee_no', 'employee_number', 'emp_id'],
            'name': ['name', 'employee_name', 'full_name'],
            'email': ['email', 'employee_email'],
            'track': ['track', 'track_name'],
            'tech_stack': ['tech_stack', 'technology', 'tech'],
            'tier': ['tier', 'tier_name'],
            'designation': ['designation', 'title', 'role'],
            'employee_type': ['employee_type', 'employment_type', 'type'],
            'university': ['university', 'uni'],
            'joined_date': ['joined_date', 'join_date', 'start_date', 'date_joined'],
            'date_of_birth': ['date_of_birth', 'dob', 'birthday'],
            'last_increment_date': ['last_increment_date', 'increment_date'],
            'last_promotion_date': ['last_promotion_date', 'promotion_date'],
            'status': ['status', 'employee_status'],
            'is_account_manager': ['is_account_manager', 'account_manager'],
            'is_external': ['is_external', 'external'],
            'global_employee_id': ['global_employee_id', 'global_id'],
            'phone_number': ['phone_number', 'phone', 'mobile'],
            'nic_passport': ['nic_passport', 'nic', 'passport'],
            'skills': ['skills', 'skill_set'],
        }

        def find_col(field_name):
            aliases = HEADER_MAP.get(field_name, [field_name])
            for alias in aliases:
                if alias in headers:
                    return headers.index(alias)
            return None

        col_map = {field: find_col(field) for field in HEADER_MAP}

        if col_map['emp_no'] is None and col_map['epf_no'] is None:
            return error("Required column 'Emp No' or 'EPF No' not found in Excel headers", 400,
                        {'found_headers': headers})

        # Pre-load lookup mappings
        TRACK_MAP = {
            'QA': 1, 'Dev': 2, 'UI': 3, 'BA': 4, 'PM': 5, 'Support': 6,
            'UX': 8, 'Execs': 9, 'Delivery': 10,
            'Functional Consultant - MS Dynamics 365': 11,
        }

        TECH_STACK_MAP = {
            'QA': 1, '.NET': 2, 'Full Stack': 3, 'Synergy': 4, 'PM': 5,
            'BA': 6, 'UI': 7, 'Java': 8, 'Data Science': 9, 'Power Apps': 10,
            'Finance': 11, 'React': 12, 'Dynamics': 13, 'UX': 14,
            'BA/PM': 15, 'UI/UX': 16, 'HR': 17, 'Execs': 18, 'Admin': 19,
            'Marketing': 20, 'Drupal': 21, 'Sales & Marketing': 22,
            'BC': 23, 'Business Central (Functional)': 24, 'AI/ML': 25,
            'Blockchain': 26,
        }

        TIER_MAP = {
            'Tier - 1': 1, 'Tier-1': 1, 'Tier 1': 1,
            'Tier - 2': 2, 'Tier-2': 2, 'Tier 2': 2,
            'Tier - 3': 3, 'Tier-3': 3, 'Tier 3': 3,
            'Tier - 4': 4, 'Tier-4': 4, 'Tier 4': 4,
            'Intern': 5, 'None': 6, 'Synergy': 7,
        }

        try:
            for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
                try:
                    def get_val(field):
                        idx = col_map.get(field)
                        return row[idx] if idx is not None and idx < len(row) else None

                    emp_no = _clean_str(get_val('emp_no'))
                    epf_no = _clean_str(get_val('epf_no'))
                    name = _clean_str(get_val('name'))

                    if not emp_no and not epf_no:
                        continue
                    if not name:
                        continue

                    email = _clean_str(get_val('email'))
                    track_name = _clean_str(get_val('track'))
                    tech_stack_name = _clean_str(get_val('tech_stack'))
                    tier_name = _clean_str(get_val('tier'))
                    designation_name = _clean_str(get_val('designation'))
                    emp_type_name = _clean_str(get_val('employee_type'))
                    university_name = _clean_str(get_val('university'))
                    joined_date = _parse_date(get_val('joined_date'))
                    dob = _parse_date(get_val('date_of_birth'))
                    last_increment = _parse_date(get_val('last_increment_date'))
                    last_promotion = _parse_date(get_val('last_promotion_date'))
                    status = _clean_str(get_val('status')) or 'Active'
                    is_am = str(get_val('is_account_manager')).lower() in ('true', 'yes', '1') if get_val('is_account_manager') else False
                    is_ext = str(get_val('is_external')).lower() in ('true', 'yes', '1') if get_val('is_external') else False
                    global_emp_id = _clean_str(get_val('global_employee_id'))
                    phone = _clean_str(get_val('phone_number'))
                    nic = _clean_str(get_val('nic_passport'))

                    # Resolve IDs
                    track_id = TRACK_MAP.get(track_name)
                    tech_stack_id = TECH_STACK_MAP.get(tech_stack_name)
                    tier_id = TIER_MAP.get(tier_name)

                    designation_id = None
                    if designation_name:
                        cursor.execute("SELECT id FROM designations WHERE name = %s", (designation_name,))
                        d_row = cursor.fetchone()
                        if d_row:
                            designation_id = d_row['id']

                    emp_type_id = None
                    if emp_type_name:
                        cursor.execute("SELECT id FROM employee_types WHERE name = %s", (emp_type_name,))
                        et_row = cursor.fetchone()
                        if et_row:
                            emp_type_id = et_row['id']

                    university_id = None
                    if university_name:
                        cursor.execute("SELECT id FROM universities WHERE name = %s", (university_name,))
                        u_row = cursor.fetchone()
                        if u_row:
                            university_id = u_row['id']
                        else:
                            cursor.execute(
                                "INSERT INTO universities (name, is_active) VALUES (%s, TRUE) RETURNING id",
                                (university_name,)
                            )
                            university_id = cursor.fetchone()['id']

                    # Upsert employee by emp_no
                    cursor.execute("""
                        INSERT INTO employees (
                            epf_no, emp_no, name, email, track_id, tech_stack_id, tier_id,
                            designation_id, employee_type_id, university_id,
                            joined_date, date_of_birth, last_increment_date, last_promotion_date,
                            status, is_account_manager, is_external,
                            global_employee_id, phone_number, nic_passport
                        ) VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
                        ON CONFLICT (emp_no) WHERE deleted_at IS NULL
                        DO UPDATE SET
                            epf_no = COALESCE(EXCLUDED.epf_no, employees.epf_no),
                            name = EXCLUDED.name,
                            email = COALESCE(EXCLUDED.email, employees.email),
                            track_id = COALESCE(EXCLUDED.track_id, employees.track_id),
                            tech_stack_id = COALESCE(EXCLUDED.tech_stack_id, employees.tech_stack_id),
                            tier_id = COALESCE(EXCLUDED.tier_id, employees.tier_id),
                            designation_id = COALESCE(EXCLUDED.designation_id, employees.designation_id),
                            employee_type_id = COALESCE(EXCLUDED.employee_type_id, employees.employee_type_id),
                            university_id = COALESCE(EXCLUDED.university_id, employees.university_id),
                            joined_date = COALESCE(EXCLUDED.joined_date, employees.joined_date),
                            date_of_birth = COALESCE(EXCLUDED.date_of_birth, employees.date_of_birth),
                            last_increment_date = COALESCE(EXCLUDED.last_increment_date, employees.last_increment_date),
                            last_promotion_date = COALESCE(EXCLUDED.last_promotion_date, employees.last_promotion_date),
                            status = EXCLUDED.status,
                            is_account_manager = EXCLUDED.is_account_manager,
                            is_external = EXCLUDED.is_external,
                            global_employee_id = COALESCE(EXCLUDED.global_employee_id, employees.global_employee_id),
                            phone_number = COALESCE(EXCLUDED.phone_number, employees.phone_number),
                            nic_passport = COALESCE(EXCLUDED.nic_passport, employees.nic_passport),
                            updated_at = NOW()
                        RETURNING (xmax = 0) AS is_insert
                    """, (
                        epf_no, emp_no, name, email, track_id, tech_stack_id, tier_id,
                        designation_id, emp_type_id, university_id,
                        joined_date, dob, last_increment, last_promotion,
                        status, is_am, is_ext,
                        global_emp_id, phone, nic
                    ))

                    result_row = cursor.fetchone()
                    if result_row and result_row['is_insert']:
                        results['imported'] += 1
                    else:
                        results['updated'] += 1

                except Exception as row_err:
                    results['failed'] += 1
                    results['errors'].append({'row': row_idx, 'error': str(row_err)})
                    logger.warning(f"Row {row_idx} failed: {str(row_err)}")

            conn.commit()
            logger.info(f"Employee upload complete: {results}")
            return success({
                'message': 'Employee upload completed',
                'results': results
            })

        except Exception as e:
            conn.rollback()
            raise e

    except ValueError as ve:
        return error(str(ve), 400)
    except Exception as e:
        logger.error(f"Employee upload failed: {str(e)}", exc_info=True)
        return error(f"Employee upload failed: {str(e)}", 500)


# ============================================================================
# ALLOCATION UPLOAD
# ============================================================================

def upload_allocations(event, context):
    """
    Upload allocations from Excel file.
    Expected columns: Emp No, Employee Name, Project ID/Code, Project Name,
                     Project Assign Date, Allocated Role, Billing Status,
                     Resource Billing, Project Allocation, Project Deassigned Date,
                     Note, Status, Global Employee ID, Resource Cost
    """
    logger.info("Starting allocation upload from Excel")

    try:
        file_buffer = _parse_excel_from_event(event)
        wb = openpyxl.load_workbook(file_buffer, read_only=True, data_only=True)
        ws = wb.active

        headers = _get_headers(ws)
        logger.info(f"Found headers: {headers}")

        conn = get_connection()
        conn.autocommit = False
        cursor = conn.cursor()

        results = {'imported': 0, 'updated': 0, 'failed': 0, 'skipped': 0, 'errors': []}

        HEADER_MAP = {
            'emp_no': ['emp_no', 'employee_no', 'employee_number', 'emp_id'],
            'employee_name': ['employee_name', 'name', 'resource_name'],
            'project_code': ['project_code', 'project_id', 'project_no'],
            'project_name': ['project_name', 'project'],
            'allocated_date': ['allocated_date', 'project_assign_date', 'assign_date', 'start_date'],
            'allocated_role': ['allocated_role', 'role', 'position'],
            'billing_status': ['billing_status', 'billing'],
            'resource_billing': ['resource_billing', 'billing_percentage', 'billing_%'],
            'allocation_percentage': ['allocation_percentage', 'project_allocation', 'allocation', 'allocation_%'],
            'deallocated_date': ['deallocated_date', 'project_deassigned_date', 'deassign_date', 'end_date'],
            'notes': ['notes', 'note', 'comments'],
            'status': ['status'],
            'global_employee_id': ['global_employee_id', 'global_id'],
            'resource_cost': ['resource_cost', 'cost'],
        }

        def find_col(field_name):
            aliases = HEADER_MAP.get(field_name, [field_name])
            for alias in aliases:
                if alias in headers:
                    return headers.index(alias)
            return None

        col_map = {field: find_col(field) for field in HEADER_MAP}

        if col_map['emp_no'] is None and col_map['employee_name'] is None:
            return error("Required column 'Emp No' or 'Employee Name' not found", 400,
                        {'found_headers': headers})

        try:
            for row_idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
                try:
                    def get_val(field):
                        idx = col_map.get(field)
                        return row[idx] if idx is not None and idx < len(row) else None

                    emp_no = _clean_str(get_val('emp_no'))
                    emp_name = _clean_str(get_val('employee_name'))
                    project_code = _clean_str(get_val('project_code'))
                    project_name = _clean_str(get_val('project_name'))

                    if not emp_no and not emp_name:
                        continue

                    # Resolve employee_id
                    employee_id = None
                    if emp_no:
                        cursor.execute("SELECT id FROM employees WHERE emp_no = %s AND deleted_at IS NULL", (emp_no,))
                        emp_row = cursor.fetchone()
                        if emp_row:
                            employee_id = emp_row['id']
                    if not employee_id and emp_name:
                        cursor.execute(
                            "SELECT id FROM employees WHERE name ILIKE %s AND deleted_at IS NULL LIMIT 1",
                            (f"%{emp_name}%",)
                        )
                        emp_row = cursor.fetchone()
                        if emp_row:
                            employee_id = emp_row['id']

                    if not employee_id:
                        results['skipped'] += 1
                        results['errors'].append({
                            'row': row_idx,
                            'error': f"Employee not found: {emp_no or emp_name}"
                        })
                        continue

                    # Resolve project_id
                    project_id = None
                    if project_code:
                        cursor.execute(
                            "SELECT id FROM projects WHERE project_code = %s AND deleted_at IS NULL",
                            (project_code,)
                        )
                        p_row = cursor.fetchone()
                        if p_row:
                            project_id = p_row['id']
                    if not project_id and project_name:
                        cursor.execute(
                            "SELECT id FROM projects WHERE project_name = %s AND deleted_at IS NULL",
                            (project_name,)
                        )
                        p_row = cursor.fetchone()
                        if p_row:
                            project_id = p_row['id']

                    if not project_id:
                        results['skipped'] += 1
                        results['errors'].append({
                            'row': row_idx,
                            'error': f"Project not found: {project_code or project_name}"
                        })
                        continue

                    allocated_date = _parse_date(get_val('allocated_date'))
                    deallocated_date = _parse_date(get_val('deallocated_date'))
                    allocation_pct = _clean_int(get_val('allocation_percentage')) or 100
                    billing_pct = _clean_int(get_val('resource_billing')) or 0
                    notes = _clean_str(get_val('notes'))

                    # Resolve billing_status_id
                    billing_status_name = _clean_str(get_val('billing_status'))
                    billing_status_id = None
                    if billing_status_name:
                        cursor.execute("SELECT id FROM billing_statuses WHERE name = %s", (billing_status_name,))
                        bs_row = cursor.fetchone()
                        if bs_row:
                            billing_status_id = bs_row['id']

                    # Get system user
                    cursor.execute("SELECT id FROM users WHERE username = 'system' LIMIT 1")
                    system_user = cursor.fetchone()
                    created_by = system_user['id'] if system_user else 1

                    is_active = True
                    status_val = _clean_str(get_val('status'))
                    if status_val and status_val.lower() in ('inactive', 'deallocated', 'ended', 'closed'):
                        is_active = False

                    # Upsert allocation
                    cursor.execute("""
                        INSERT INTO allocations (
                            employee_id, project_id, allocation_percentage, billing_percentage,
                            billing_status_id, allocated_date, deallocated_date,
                            effective_date, is_active, notes, created_by
                        ) VALUES (%s, %s, %s, %s, %s, %s, %s, COALESCE(%s, CURRENT_DATE), %s, %s, %s)
                        ON CONFLICT (employee_id, project_id) WHERE deleted_at IS NULL AND is_active = true
                        DO UPDATE SET
                            allocation_percentage = EXCLUDED.allocation_percentage,
                            billing_percentage = EXCLUDED.billing_percentage,
                            billing_status_id = COALESCE(EXCLUDED.billing_status_id, allocations.billing_status_id),
                            allocated_date = COALESCE(EXCLUDED.allocated_date, allocations.allocated_date),
                            deallocated_date = EXCLUDED.deallocated_date,
                            is_active = EXCLUDED.is_active,
                            notes = COALESCE(EXCLUDED.notes, allocations.notes),
                            updated_at = NOW()
                        RETURNING (xmax = 0) AS is_insert
                    """, (
                        employee_id, project_id, allocation_pct, billing_pct,
                        billing_status_id, allocated_date, deallocated_date,
                        allocated_date, is_active, notes, created_by
                    ))

                    result_row = cursor.fetchone()
                    if result_row and result_row['is_insert']:
                        results['imported'] += 1
                    else:
                        results['updated'] += 1

                except Exception as row_err:
                    results['failed'] += 1
                    results['errors'].append({'row': row_idx, 'error': str(row_err)})
                    logger.warning(f"Row {row_idx} failed: {str(row_err)}")

            conn.commit()
            logger.info(f"Allocation upload complete: {results}")
            return success({
                'message': 'Allocation upload completed',
                'results': results
            })

        except Exception as e:
            conn.rollback()
            raise e

    except ValueError as ve:
        return error(str(ve), 400)
    except Exception as e:
        logger.error(f"Allocation upload failed: {str(e)}", exc_info=True)
        return error(f"Allocation upload failed: {str(e)}", 500)
